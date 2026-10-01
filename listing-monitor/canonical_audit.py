"""Site-wide audit: 26/27 listings whose canonical points to another page.

A product page whose Rank Math canonical points to another URL (usually the
listing it was duplicated from, e.g. last season's) is not indexed by Google.

Usage:
  python canonical_audit.py                 # all sites
  python canonical_audit.py --sites KFK RFS

Output: reports/canonical_<date>/canonical_26-27_<date>.xlsx (one sheet per site)
"""

import argparse
import html
import json
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from check_new_listings import BASE_DIR, REPORT_DIR, get_tz, http_get
from qa_fetch import fetch_head, parse_head

SEASON_2627 = re.compile(r"(?<!\d)(?:20)?26\s*[/_-]\s*27(?!\d)")
OLD_SEASON = re.compile(r"(?<!\d)(?:20)?(?:1\d|2[0-5])\s*[-_/]\s*\d{2}(?!\d)|(?<!\d)20(?:1\d|2[0-5])(?!\d)")


def list_products(base_url):
    """id, name, link of every published product (100 per request)."""
    out, page = [], 1
    while True:
        text, headers = http_get(f"{base_url}/wp-json/wp/v2/product?per_page=100&page={page}"
                                 f"&_fields=id,title,link,date", retry_waits=(20, 60))
        batch = json.loads(text)
        out += [{"id": p["id"], "name": html.unescape(p["title"]["rendered"]), "url": p["link"],
                 "date": p["date"][:10]} for p in batch]
        if page >= int(headers.get("X-WP-TotalPages") or page) or not batch:
            return out
        page += 1


def norm(u):
    return (u or "").split("#")[0].split("?")[0].rstrip("/").lower()


def audit_site(site, base_url, delay, progress):
    products = [p for p in list_products(base_url) if SEASON_2627.search(p["name"])]
    # Progress is saved after every page so a stopped run resumes where it left off (same day only).
    done_path = BASE_DIR / "data" / "cache" / f"canonical_{site}_{datetime.now().date().isoformat()}.json"
    done = json.loads(done_path.read_text(encoding="utf-8")) if done_path.exists() else {}
    done_path.parent.mkdir(parents=True, exist_ok=True)
    rows, errors = [], []
    for n, p in enumerate(products, 1):
        page = done.get(p["url"])
        if page is None:
            try:
                page = parse_head(fetch_head(p["url"]))
            except Exception as e:
                errors.append((p, str(e)))
                continue
            done[p["url"]] = page
            done_path.write_text(json.dumps(done), encoding="utf-8")
            time.sleep(delay)
        canon = page["canonical"]
        if canon and norm(canon) != norm(p["url"]):
            slug = canon.rstrip("/").rsplit("/", 1)[-1]
            rows.append({**p, "canonical": canon,
                         "old_season": bool(OLD_SEASON.search(slug.replace("2026", ""))),
                         "robots": page["robots"]})
        if n % 50 == 0:
            progress(f"{site}: {n}/{len(products)}")
    progress(f"{site}: xong {len(products)} listing 26/27, {len(rows)} canonical sai, {len(errors)} loi tai")
    return site, len(products), rows, errors


def write_excel(path, results, when):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill, head_font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
    red = PatternFill("solid", fgColor="F8D7DA")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    ws.append([f"Canonical sai trên listing 26/27 (check lúc {when})"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Canonical trỏ sang trang khác -> Google coi listing là bản trùng và KHÔNG index. "
               "Sửa: Rank Math -> Advanced -> xoá ô Canonical URL."])
    ws.append([])
    ws.append(["Site", "Listing 26/27 đã check", "Canonical sai", "Trong đó trỏ về trang mùa cũ", "Không tải được"])
    for c in ws[ws.max_row]:
        c.fill, c.font = head_fill, head_font
    for site, total, rows, errors in results:
        ws.append([site, total, len(rows), sum(r["old_season"] for r in rows), len(errors)])
    for col, w in zip("ABCDE", (10, 22, 14, 28, 16)):
        ws.column_dimensions[col].width = w

    for site, total, rows, errors in results:
        ws = wb.create_sheet(site)
        ws.append(["ID", "Ngày đăng", "Tên sản phẩm", "URL listing", "Canonical đang trỏ về",
                   "Trỏ về mùa cũ?", "Sửa trong admin", "Người sửa", "Đã sửa"])
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for p in sorted(rows, key=lambda r: (not r["old_season"], r["date"])):
            base = p["url"].split("/product/")[0]
            ws.append([p["id"], p["date"], p["name"], p["url"], p["canonical"],
                       "Có" if p["old_season"] else "", "Mở", "", ""])
            r = ws.max_row
            ws.cell(r, 4).hyperlink = p["url"]
            ws.cell(r, 5).hyperlink = p["canonical"]
            ws.cell(r, 7).hyperlink = f"{base}/wp-admin/post.php?post={p['id']}&action=edit"
            ws.cell(r, 7).font = Font(color="0563C1", underline="single")
            if p["old_season"]:
                ws.cell(r, 6).fill = red
        for p, e in errors:
            ws.append([p["id"], p["date"], p["name"], p["url"], f"KHÔNG TẢI ĐƯỢC: {e}"])
        for col, w in zip("ABCDEFGHI", (9, 12, 60, 50, 60, 14, 14, 14, 10)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            row[2].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--sites", nargs="*")
    ap.add_argument("--config", default=str(BASE_DIR / "config.json"))
    args = ap.parse_args(argv)
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    delays = cfg.get("request_delay", {})
    sites = {k: v.rstrip("/") for k, v in cfg["sites"].items() if not args.sites or k in args.sites}
    now = datetime.now(get_tz(cfg.get("timezone", "Europe/London")))

    with ThreadPoolExecutor(len(sites)) as ex:
        futures = [ex.submit(audit_site, s, u, delays.get(s, delays.get("default", 0.3)),
                             lambda m: print(m, flush=True)) for s, u in sites.items()]
        results = [f.result() for f in futures]

    out_dir = REPORT_DIR / f"canonical_{now.date().isoformat()}"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"canonical_26-27_{now.date().isoformat()}.xlsx"
    write_excel(path, results, now.strftime("%Y-%m-%d %H:%M"))
    for site, total, rows, errors in results:
        print(f"{site}: {total} listing 26/27, {len(rows)} canonical sai "
              f"({sum(r['old_season'] for r in rows)} tro ve mua cu), {len(errors)} khong tai duoc")
    print(f"File Excel: {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
