"""Deep QA of new listings -> Excel report.

Reads the new-listing CSV made by check_new_listings.py, downloads each
listing (qa_fetch), runs the rules (qa_rules) and writes:
  reports/<label>/qa_<label>.xlsx   (Tong quan / Loi chi tiet / Can xac nhan)
  reports/<label>/qa_summary.txt    (short text for the daily message)

Usage:
  python qa_listings.py                                  # today
  python qa_listings.py --date 2026-09-25
  python qa_listings.py --from 2026-08-26 --to 2026-09-26
"""

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from pathlib import Path

from check_new_listings import BASE_DIR, REPORT_DIR, get_tz
import qa_fetch
import qa_rules

SEV_ORDER = {qa_rules.HIGH: 0, qa_rules.MED: 1, qa_rules.LOW: 2}
SEV_LABEL = {qa_rules.HIGH: "Cao", qa_rules.MED: "Trung bình", qa_rules.LOW: "Thấp"}
SEV_FILL = {qa_rules.HIGH: "F8D7DA", qa_rules.MED: "FFF3CD", qa_rules.LOW: "E2E3E5"}
SITE_ORDER = ["KFK", "RFS", "RFK", "CFS"]
# Same problem on at least this share of a site's listings -> probably the site template.
TEMPLATE_SHARE = 0.5
TEMPLATE_MIN = 5

TO_CONFIRM = [
    ("SKU ADK", "SOP ghi ADK = combo Adult + Kid, nhưng trên web ADK đang dùng cho Adult Kit (áo + quần người lớn). "
                "Tool đang hiểu ADK = Adult. Đúng không?"),
    ("Short sleeve", "SOP: không bao giờ nhắc short sleeve. Nhưng template description KFK có 'Short-Sleeve shirt' và "
                     "'Sleeve length: Short sleeve'. Có cần xoá khỏi template không?"),
    ("Trustpilot / 4.9★", "SOP: tránh 'Trustpilot' và claim 4.9★. Short description KFK và một số meta đang dùng. "
                          "Đây là template được duyệt hay cần sửa?"),
    ("Tên né thương hiệu", "Tool coi 'North London Red' = Arsenal, 'P^_^SG' = PSG. Còn tên thay thế nào khác không?"),
    ("WINNERS / CHAMPIONS", "Listing in chữ kỷ niệm (vd 'CHAMPIONS 26') không bị báo thiếu player attribute/category. Đúng không?"),
    ("Gift pack NOINDEX", "Một số gift pack KFK (Bay^_^rn, P^_^SG) đang để noindex. Cố ý (né thương hiệu) hay lỗi?"),
    ("Sleeve badge", "Tool KHÔNG báo lỗi 'sleeve badge' (tên option), chỉ báo 'badge' đứng riêng. Đúng không?"),
    ("Độ dài SEO title", "Tool báo khi SEO title > 65 ký tự và meta > 165 ký tự (mức Thấp)."),
    ("KFK / CFS / RFK template SEO", "SOP chưa có template SEO title/meta cho KFK, CFS, RFK nên tool chưa check đúng mẫu "
                                     "cho các site này, chỉ check nội dung có khớp sản phẩm không."),
]


def load_listings(label):
    path = REPORT_DIR / label / f"new_listings_{label}.csv"
    if not path.exists():
        sys.exit(f"Chua co {path}. Chay check_new_listings.py truoc.")
    with path.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def fetch_all(listings, cfg, refresh):
    """One thread per site so no single site gets a burst of requests."""
    delays = cfg.get("request_delay", {})
    by_site = defaultdict(list)
    for r in listings:
        by_site[r["site"]].append(r)
    products, errors = [], []

    def run(site):
        out = []
        base = cfg["sites"][site].rstrip("/")
        todo = [int(r["product_id"]) for r in by_site[site]
                if refresh or not qa_fetch.cached(site, int(r["product_id"]))]
        try:
            store = qa_fetch.prefetch_store(base, todo) if todo else {}
        except Exception:
            store = {}  # fall back to one request per product
        for r in by_site[site]:
            try:
                d = qa_fetch.fetch_product(site, base, int(r["product_id"]), r["url"],
                                           refresh=refresh, delay=delays.get(site, delays.get("default", 0.3)),
                                           store=store.get(int(r["product_id"])))
                d["published_at"] = r["published_at"]
                out.append(d)
            except Exception as e:
                errors.append((site, r["product_id"], r["name"], str(e)))
        return out

    with ThreadPoolExecutor(len(by_site) or 1) as ex:
        for out in ex.map(run, list(by_site)):
            products.extend(out)
    return products, errors


def run_checks(products):
    rows = []
    for d in products:
        for i in qa_rules.check_product(d):
            rows.append((d, i))
    rows.extend(qa_rules.check_batch(products))
    return rows


def issue_key(i):
    """Groups the same rule across products (detail minus the product-specific part)."""
    head = i["detail"].split(":")[0].split("(")[0]
    return (i["group"], i["field"], head[:60].strip())


def find_template_issues(products, rows):
    per_site = Counter(d["site"] for d in products)
    hit = defaultdict(set)
    for d, i in rows:
        hit[(d["site"],) + issue_key(i)].add(d["store"]["id"])
    out = {}
    for k, ids in hit.items():
        n = per_site[k[0]]
        if len(ids) >= TEMPLATE_MIN and len(ids) / n >= TEMPLATE_SHARE:
            out[k] = (len(ids), n)
    return out


def write_excel(path, label, products, rows, errors, templates):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, cols, widths):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.fill, c.font = head_fill, head_font
        for n, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(n)].width = w

    wb = Workbook()

    # ---- Tong quan
    ws = wb.active
    ws.title = "Tổng quan"
    ws.append([f"QA listing mới: {label.replace('_to_', ' → ')}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Tạo lúc {datetime.now().strftime('%Y-%m-%d %H:%M')}. Mức độ: Cao = sai thông tin / gây confuse "
               "khách và Google; Trung bình = cần sửa; Thấp = nên sửa."])
    ws.append([])
    sites = [s for s in SITE_ORDER if any(d["site"] == s for d in products)]
    by_site_prod = Counter(d["site"] for d in products)
    prod_high = defaultdict(set)
    prod_any = defaultdict(set)
    sev_count = Counter()
    for d, i in rows:
        prod_any[d["site"]].add(d["store"]["id"])
        if i["severity"] == qa_rules.HIGH:
            prod_high[d["site"]].add(d["store"]["id"])
        sev_count[(d["site"], i["severity"])] += 1
    header(ws, ["Site", "Listing đã check", "Listing có lỗi", "Listing có lỗi Cao",
                "Lỗi Cao", "Lỗi Trung bình", "Lỗi Thấp"], [34, 16, 14, 18, 10, 14, 10])
    for s in sites:
        ws.append([s, by_site_prod[s], len(prod_any[s]), len(prod_high[s]),
                   sev_count[(s, qa_rules.HIGH)], sev_count[(s, qa_rules.MED)], sev_count[(s, qa_rules.LOW)]])
    ws.append([])

    ws.append(["Số listing bị lỗi theo nhóm"])
    ws.cell(ws.max_row, 1).font = bold
    ws.append(["Nhóm lỗi"] + sites + ["Tổng"])
    for c in ws[ws.max_row]:
        c.fill, c.font = head_fill, head_font
    grp = defaultdict(set)
    grp_sev = {}
    for d, i in rows:
        grp[(i["group"], d["site"])].add(d["store"]["id"])
        grp_sev[i["group"]] = min(grp_sev.get(i["group"], 9), SEV_ORDER[i["severity"]])
    groups = sorted({g for g, _ in grp}, key=lambda g: (grp_sev[g], g))
    for g in groups:
        vals = [len(grp[(g, s)]) for s in sites]
        ws.append([g] + vals + [sum(vals)])
    ws.append([])

    if templates:
        ws.append(["Lỗi lặp lại trên phần lớn listing của 1 site (nhiều khả năng do TEMPLATE, sửa template 1 lần)"])
        ws.cell(ws.max_row, 1).font = bold
        ws.append(["Site", "Nhóm", "Vị trí", "Lỗi", "Số listing"])
        for c in ws[ws.max_row]:
            c.fill, c.font = head_fill, head_font
        for (s, g, f, h), (n, tot) in sorted(templates.items()):
            ws.append([s, g, f, h, f"{n}/{tot}"])
        ws.append([])

    if errors:
        ws.append(["Không tải được (chưa check)"])
        ws.cell(ws.max_row, 1).font = bold
        for e in errors:
            ws.append(list(e))

    # ---- Loi chi tiet
    ws = wb.create_sheet("Lỗi chi tiết")
    cols = ["Mức độ", "Site", "Ngày đăng", "ID", "SKU", "Tên sản phẩm", "Nhóm lỗi", "Vị trí",
            "Chi tiết", "Cách sửa", "Lỗi template?", "URL", "Sửa trong admin"]
    header(ws, cols, [11, 6, 16, 9, 30, 45, 20, 20, 70, 40, 12, 40, 20])
    tmpl_keys = {k for k in templates}
    ordered = sorted(rows, key=lambda r: (SEV_ORDER[r[1]["severity"]], SITE_ORDER.index(r[0]["site"]),
                                          r[0].get("published_at", ""), r[0]["store"]["id"]))
    for d, i in ordered:
        s = d["store"]
        base = s["permalink"].split("/product/")[0]
        edit = f"{base}/wp-admin/post.php?post={s['id']}&action=edit"
        ws.append([SEV_LABEL[i["severity"]], d["site"], d.get("published_at", ""), s["id"], s.get("sku", ""),
                   qa_rules.strip_html(s["name"]), i["group"], i["field"], i["detail"], i["fix"],
                   "Có" if (d["site"],) + issue_key(i) in tmpl_keys else "",
                   s["permalink"], "Mở"])
        r = ws.max_row
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=SEV_FILL[i["severity"]])
        ws.cell(r, 12).hyperlink = s["permalink"]
        ws.cell(r, 13).hyperlink = edit
        ws.cell(r, 13).font = Font(color="0563C1", underline="single")
        for c in (9, 10):
            ws.cell(r, c).alignment = wrap
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---- Can xac nhan
    ws = wb.create_sheet("Cần xác nhận")
    header(ws, ["Quy tắc", "Câu hỏi cho Clara"], [24, 110])
    for a, b in TO_CONFIRM:
        ws.append([a, b])
        ws.cell(ws.max_row, 2).alignment = wrap

    wb.save(path)


def build_summary(label, products, rows, errors, templates):
    lines = [f"QA LISTING MOI {label.replace('_to_', ' -> ')}", ""]
    for s in SITE_ORDER:
        ps = [d for d in products if d["site"] == s]
        if not ps:
            continue
        rs = [(d, i) for d, i in rows if d["site"] == s]
        high = {d["store"]["id"] for d, i in rs if i["severity"] == qa_rules.HIGH}
        lines.append(f"{s}: {len(ps)} listing, {len(high)} listing co loi CAO, "
                     f"{len({d['store']['id'] for d, _ in rs})} listing co loi")
        top = Counter(i["group"] for d, i in rs if i["severity"] == qa_rules.HIGH
                      and (s,) + issue_key(i) not in templates)
        for g, n in top.most_common(5):
            lines.append(f"   - {g}: {n} loi")
    if templates:
        lines += ["", "Loi lap lai (co the do template):"]
        for (s, g, f, h), (n, tot) in sorted(templates.items()):
            lines.append(f"   - {s} | {g} | {f}: {h} ({n}/{tot})")
    if errors:
        lines += ["", f"Khong tai duoc {len(errors)} listing (xem sheet Tong quan)"]
    return "\n".join(lines)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Deep QA of new listings.")
    ap.add_argument("--date")
    ap.add_argument("--from", dest="date_from")
    ap.add_argument("--to", dest="date_to")
    ap.add_argument("--refresh", action="store_true", help="download again, ignore cache")
    ap.add_argument("--config", default=str(BASE_DIR / "config.json"))
    args = ap.parse_args(argv)

    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    today = datetime.now(get_tz(cfg.get("timezone", "Europe/London"))).date()
    if args.date_from:
        start = date.fromisoformat(args.date_from)
        end = date.fromisoformat(args.date_to) if args.date_to else today
    else:
        start = end = date.fromisoformat(args.date) if args.date else today
    label = start.isoformat() if start == end else f"{start.isoformat()}_to_{end.isoformat()}"

    listings = load_listings(label)
    products, errors = fetch_all(listings, cfg, args.refresh)
    rows = run_checks(products)
    templates = find_template_issues(products, rows)

    out = REPORT_DIR / label / f"qa_{label}.xlsx"
    write_excel(out, label, products, rows, errors, templates)
    summary = build_summary(label, products, rows, errors, templates)
    (REPORT_DIR / label / "qa_summary.txt").write_text(summary, encoding="utf-8")
    print(summary)
    print(f"\nFile Excel: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
